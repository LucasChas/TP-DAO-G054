"""Genera la documentación de la auditoría a partir de una única fuente de contenido.

Salidas:
  e2e/features/*.feature            casos en Gherkin (español)
  docs/pruebas/casos_xray.csv       planilla para importar en Jira/Xray (Cucumber)
  docs/informe_auditoria.html       documento integral navegable
  docs/Presentacion-grupo-G054.docx el mismo contenido sobre la plantilla de la cátedra

Ejecutar desde la raíz del repositorio, después de correr las pruebas E2E con --junitxml:
  PYTHONPATH=. python docs/auditoria/construir_documentacion.py
"""
import csv
import html
import re
import sys
import xml.etree.ElementTree as ET
from collections import Counter, OrderedDict
from datetime import date
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[2]
DOCS = RAIZ / "docs"
sys.path.insert(0, str(RAIZ))
sys.path.insert(0, str(Path(__file__).parent))

from e2e.casos import CASOS  # noqa: E402
from requisitos import HALLAZGOS, REQUISITOS  # noqa: E402

HOY = date.today().strftime("%d/%m/%Y")
INTEGRANTES = [
    ("95346", "Chas Diaz Lucas Martin"),
    ("408708", "Melina Micaela Moran Gianotti"),
    ("403112", "Fabrizio Sebastián Banffi Moreno"),
    ("94488", "Lorenzo Francisco Gonzalez"),
    ("94183", "Graña Carranza Mateo"),
]
REPO = "https://github.com/LucasChas/TP-DAO-G054/blob/claude/eager-wright-hnv32r/"

# --------------------------------------------------------------------------------- datos de entrada


def resultados_junit():
    """Nombre de la prueba -> (estado, segundos) según el XML de pytest."""
    archivo = DOCS / "pruebas" / "resultados_e2e.xml"
    resultados = {}
    for caso in ET.parse(archivo).iter("testcase"):
        nombre = caso.get("name")
        if caso.find("failure") is not None or caso.find("error") is not None:
            estado = "Falla"
        elif (s := caso.find("skipped")) is not None:
            estado = "Falla esperada (defecto)" if "xfail" in (s.get("type", "") + s.get("message", "")) else "Omitida"
        else:
            estado = "Pasa"
        resultados[nombre] = (estado, float(caso.get("time", 0)))
    return resultados


def resultado_unitarios():
    texto = (DOCS / "pruebas" / "resultados_unitarios.txt").read_text(encoding="utf-8")
    return texto.strip().splitlines()[-1].strip("= ")


def capturas_de(caso_id):
    return sorted((DOCS / "capturas").glob(f"{caso_id}_*.png"))


def fragmento(ruta, desde, hasta):
    lineas = (RAIZ / ruta).read_text(encoding="utf-8").splitlines()
    return "\n".join(lineas[desde - 1:hasta])


def verificar_referencias():
    """Cada archivo:línea citado tiene que existir. Devuelve la lista de referencias rotas."""
    rotas = []
    citas = [d for r in REQUISITOS for d in r.donde] + [e for h in HALLAZGOS for e in h.evidencia]
    for cita in citas:
        m = re.match(r"([\w./\-]+?)(?::(\d+)(?:-(\d+))?)?(?:\s|$)", cita)
        ruta, linea = m.group(1), m.group(2)
        archivo = RAIZ / ruta
        if ruta == "docs/Presentacion-grupo-G054.docx":  # lo genera este mismo script
            continue
        if not archivo.exists():
            rotas.append(cita)
        elif linea and archivo.suffix in (".py", ".md") and int(linea) > len(archivo.read_text(encoding="utf-8").splitlines()):
            rotas.append(cita)
    return rotas


# --------------------------------------------------------------------------------- Gherkin y Xray


def escribir_features():
    carpeta = RAIZ / "e2e" / "features"
    carpeta.mkdir(exist_ok=True)
    por_area = OrderedDict()
    for caso in CASOS:
        por_area.setdefault(caso.funcionalidad, []).append(caso)
    archivos = {}
    for area, casos in por_area.items():
        nombre = re.sub(r"[^a-z]+", "_", area.lower().replace("á", "a").replace("é", "e").replace("í", "i")
                        .replace("ó", "o").replace("ú", "u")).strip("_") + ".feature"
        lineas = ["# language: es", f"Característica: {area}",
                  "  Pruebas E2E sobre la interfaz gráfica real (Tkinter). Implementación: e2e/test_e2e_*.py", ""]
        for c in casos:
            etiquetas = [f"@{c.id}", "@positivo" if c.tipo == "CP" else "@negativo"] + [f"@{r}" for r in c.requisitos]
            if c.marca == "xfail":
                etiquetas.append("@defecto_conocido")
            lineas.append("  " + " ".join(etiquetas))
            lineas.append(f"  Escenario: {c.id} {c.descripcion}")
            for palabra, pasos in (("Dado", c.dado), ("Cuando", c.cuando), ("Entonces", c.entonces)):
                for i, paso in enumerate(pasos):
                    lineas.append(f"    {palabra if i == 0 else 'Y'} {paso}")
            lineas.append("")
        (carpeta / nombre).write_text("\n".join(lineas), encoding="utf-8")
        archivos[area] = f"e2e/features/{nombre}"
    return archivos


def gherkin_de(caso):
    lineas = []
    for palabra, pasos in (("Given", caso.dado), ("When", caso.cuando), ("Then", caso.entonces)):
        for i, paso in enumerate(pasos):
            lineas.append(f"{palabra if i == 0 else 'And'} {paso}")
    return "\n".join(lineas)


def escribir_xray():
    ruta = DOCS / "pruebas" / "casos_xray.csv"
    with ruta.open("w", newline="", encoding="utf-8-sig") as archivo:
        w = csv.writer(archivo, delimiter=";")
        w.writerow(["Summary", "Description", "Repositorio", "Tipo de test", "Gherkin"])
        for c in CASOS:
            w.writerow([c.summary, "\n".join(f"- {x}" for x in c.criterios), f"G054-TPI2026\\{c.funcionalidad}", "Cucumber", gherkin_de(c)])
    return ruta


# --------------------------------------------------------------------------------- contenido

B = []  # bloques del documento: tuplas (tipo, ...)


def h(nivel, texto, ancla=None):
    B.append(("h", nivel, texto, ancla))


def p(texto):
    B.append(("p", texto))


def ul(items):
    B.append(("ul", list(items)))


def tabla(encabezados, filas, clase=""):
    B.append(("table", list(encabezados), [list(f) for f in filas], clase))


def img(ruta, pie, ancho_cm=16.5):
    B.append(("img", ruta, pie, ancho_cm))


def codigo(texto, origen=""):
    B.append(("code", texto, origen))


def nota(texto, tipo="info"):
    B.append(("note", texto, tipo))


def galeria(rutas):
    B.append(("gallery", list(rutas)))


def detalle_inicio(titulo, estado=""):
    B.append(("details_start", titulo, estado))


def detalle_fin():
    B.append(("details_end",))


def construir_contenido(resultados, features):
    obligatorios = [r for r in REQUISITOS if r.obligatorio]
    cuenta = Counter(r.estado for r in obligatorios)
    porcentaje = (cuenta["Cumple"] + 0.5 * cuenta["Parcial"]) / len(obligatorios) * 100
    estados_e2e = Counter(resultados.get(c.prueba.split("::")[1], ("Sin ejecutar",))[0] for c in CASOS)
    unit = resultado_unitarios()
    casos_por_req = {r.id: [c.id for c in CASOS if r.id in c.requisitos] for r in REQUISITOS}
    pasan = int(re.search(r"(\d+) passed", unit).group(1))
    fallan = int(m.group(1)) if (m := re.search(r"(\d+) failed", unit)) else 0
    resumen = dict(porcentaje=porcentaje, cuenta=cuenta, total=len(obligatorios), e2e=estados_e2e, unit=unit,
                   unit_corto=f"{pasan}/{pasan + fallan}")
    pendientes = [x for x in HALLAZGOS if not x.resuelto]
    resueltos = [x for x in HALLAZGOS if x.resuelto]

    # 1 -------------------------------------------------------------------------------------- Grupo
    h(1, "1. Grupo", "grupo")
    tabla(["Dato", "Valor"], [
        ["Número de grupo", "G054"],
        ["Tema", "Tema 1 – Biblioteca: Sistema de Préstamos Interbibliotecarios (TP Integrador 2026, DAO, ISI – UTN FRC)"],
        ["Repositorio", "github.com/LucasChas/TP-DAO-G054 (código en la rama `desarrollo`)"],
    ])
    tabla(["Legajo", "Nombre"], [list(x) for x in INTEGRANTES])

    # 2 -------------------------------------------------------------------------------------- Propuesta
    h(1, "2. Propuesta", "propuesta")
    h(2, "Objetivos del sistema")
    ul([
        "Administrar las sedes de la red (datos de contacto, horarios y baja lógica), los socios y el catálogo de libros con sus ejemplares físicos por sede.",
        "Registrar préstamos locales e interbibliotecarios, con sede de origen (dueña del ejemplar) y sede de destino (retiro).",
        "Trasladar ejemplares entre sedes con remitos numerados cuyo ciclo de vida (empaquetado, despacho, puntos de control, recepción, daño, extravío, anulación, cierre) queda registrado paso a paso.",
        "Consultar la disponibilidad de un libro en cada sede y gestionar reservas por orden de llegada.",
        "Registrar devoluciones con el estado físico del material y calcular multas con una política intercambiable.",
        "Controlar las reglas de negocio de la consigna (disponibilidad, socios inhabilitados, vencimientos, duplicados, sede del ejemplar, sedes dadas de baja).",
        "Emitir cinco reportes de gestión en pantalla y exportarlos a PDF (con gráficos) y CSV.",
    ])
    h(2, "Alcances")
    tabla(["Incluye (verificado en el código)", "No incluye"], [[
        "Interfaz gráfica Tkinter (8 pestañas) e interfaz de consola sobre los mismos servicios; persistencia en SQLite con SQLAlchemy; "
        "8 procesos obligatorios; 11 validaciones; patrones State, Strategy, Iterator y DAO; 5 reportes con exportación; datos de ejemplo; 247 tests unitarios.",
        "Notificaciones, historial de lecturas y recomendaciones (opcionales de la consigna); usuarios y permisos; exportación o impresión del remito como documento; "
        "un reporte de disponibilidad de todo el catálogo por sede.",
    ]])
    h(2, "DER (esquema real de la base)")
    p("Generado a partir de las tablas que crea `Base.metadata.create_all` (claves primarias, únicas y foráneas tal como quedan en SQLite).")
    img("auditoria/diagramas/der.png", "DER del esquema SQLite generado por SQLAlchemy", 16.5)
    h(2, "Diagrama de clases UML del dominio")
    img("auditoria/diagramas/clases.png", "Clases del paquete biblioteca/modelos y su relación con el estado del remito", 12)

    # 3 -------------------------------------------------------------------------------------- Arquitectura
    h(1, "3. Arquitectura", "arquitectura")
    p("Aplicación de escritorio en Python 3.12 con arquitectura en capas. Las dos interfaces llaman a los mismos servicios; "
      "ninguna accede a la base. Cada caso de uso abre una sesión con `obtener_sesion()`, usa los DAO y confirma o deshace todo junto.")
    img("auditoria/diagramas/arquitectura.png", "Capas y dependencias del paquete biblioteca", 15)
    tabla(["Capa / paquete", "Responsabilidad", "Contenido"], [
        ["`interfaz_grafica/`", "Ventana Tkinter: una pestaña por área, formularios modales y tablas", "aplicacion.py, componentes.py, opciones.py, 8 pestania_*.py"],
        ["`interfaz/`", "Menús de consola con rich", "menu_principal.py y un menú por área"],
        ["`servicios/`", "Casos de uso y validaciones de negocio; lanzan excepciones de dominio", "8 servicios + reglas.py, validaciones.py, asignacion_reservas.py"],
        ["`patrones/`", "State, Strategy e Iterator", "estados_remito.py, estrategias_multa.py, iterador_catalogo.py"],
        ["`reportes/`", "Cálculo de 5 reportes y exportación PDF/CSV", "un módulo por reporte, armado.py, exportar_pdf.py, exportar_csv.py, graficos.py"],
        ["`dao/`", "Acceso a datos, un DAO por entidad; reciben la sesión del servicio", "8 *_dao.py"],
        ["`modelos/`", "Entidades mapeadas con SQLAlchemy y enumeraciones", "9 clases + enumeraciones.py"],
        ["`configuracion/`", "Motor, sesión (commit/rollback), creación de tablas, datos de ejemplo", "base_de_datos.py, datos_de_ejemplo.py"],
        ["`excepciones/`", "ErrorDeNegocio y 11 excepciones derivadas", "errores.py"],
    ])
    tabla(["Aspecto", "Detalle verificado"], [
        ["Tamaño", "79 módulos, 5.654 líneas en `biblioteca/`; 247 tests unitarios en `tests/` (217 funciones, algunas parametrizadas)."],
        ["Persistencia", "SQLite (`biblioteca.db` en la carpeta de ejecución) con SQLAlchemy 2.1; claves foráneas activadas con PRAGMA en cada conexión (base_de_datos.py:15)."],
        ["Transacciones", "`obtener_sesion()` es un context manager: commit si no hubo error, rollback si lo hubo (base_de_datos.py:54)."],
        ["Errores", "Las interfaces capturan `ErrorDeNegocio` y muestran el motivo sin cerrar el programa (componentes.py:43, consola.py:100)."],
        ["Dependencias", "SQLAlchemy 2.1.4, rich 15.0.0, reportlab 5.1.0, matplotlib 3.11.2, pytest 9.1.1; Tkinter del sistema."],
        ["Ejecución", "`python main.py` (consola) o `python main.py --grafica` (GUI). Pasos exactos en el Anexo."],
    ])
    nota("Relevamiento hecho sobre la rama `desarrollo` (commit 7973887), que contiene a `interfaz-grafica`. La rama `main` solo tiene los documentos de la cátedra.")

    # 4 -------------------------------------------------------------------------------------- Matriz
    h(1, "4. Matriz de cumplimiento de la consigna", "matriz")
    p(f"Cumplimiento de lo obligatorio: **{porcentaje:.1f} %** ({cuenta['Cumple']} cumplen, {cuenta['Parcial']} parciales, "
      f"{cuenta.get('No cumple', 0)} no cumplen, sobre {len(obligatorios)} requisitos; un parcial vale 0,5). "
      "Los reportes sugeridos y los procesos opcionales se listan aparte y no entran en el porcentaje.")
    filas = []
    for r in REQUISITOS:
        evidencia = ", ".join(casos_por_req[r.id]) or "-"
        if r.unitarios:
            evidencia += "; unit: " + "; ".join(r.unitarios)
        filas.append([r.id, f"{r.descripcion} ({r.grupo})", r.estado, "\n".join(f"`{d}`" for d in r.donde) or "-", evidencia, r.observaciones or "-"])
    tabla(["Req.", "Descripción", "Estado", "Dónde (archivo:línea)", "Evidencia (casos E2E / tests)", "Observaciones"], filas, "matriz")
    h(2, "Faltantes y riesgos para la defensa (por prioridad)")
    tabla(["#", "Prioridad", "Hallazgo", "Arreglo sugerido"],
          [[x.id, x.prioridad, f"**{x.titulo}.** {x.detalle}", x.arreglo] for x in pendientes])
    if resueltos:
        p("Corregidos después de la auditoría: " + ", ".join(f"{x.id} ({x.titulo.lower()})" for x in resueltos) + ". Detalle en la sección 9.")

    # 5 -------------------------------------------------------------------------------------- Patrones
    h(1, "5. Patrones de diseño", "patrones")
    p("La consigna pide al menos dos. Se verificaron cuatro aplicados de forma explícita, más algunos usos implícitos.")
    patrones = [
        ("State: ciclo de vida del remito",
         "Un remito pasa por 9 estados y en cada uno solo algunas acciones son válidas (no se recibe lo que no salió, no se anula lo que ya viaja). "
         "Además cada transición tiene consecuencias distintas sobre el préstamo y el ejemplar según el remito sea de ida o de retorno.",
         "`EstadoRemito` (abstracta, estados_remito.py:24) implementa todas las acciones lanzando `TransicionRemitoInvalida`; "
         "cada estado concreto (`Generado` :64, `Empaquetado` :76, `EnViaje` :95, `Recibido` :143, …) redefine solo las que admite. "
         "`Remito` (remito.py:12) es el contexto: delega en `self._estado` y registra cada cambio en `MovimientoRemito` (remito.py:100). "
         "El estado se persiste por nombre y se reconstruye con `@reconstructor` (remito.py:49).",
         "Sin el patrón, cada acción sería un `if remito.estado == ...` repetido en el servicio, con reglas dispersas y difíciles de revisar. "
         "Con State, agregar un estado es agregar una clase, y una acción no permitida falla siempre igual sin tocar el estado (CN-27, CN-28, CN-30).",
         [("biblioteca/patrones/estados_remito.py", 24, 37), ("biblioteca/patrones/estados_remito.py", 76, 92),
          ("biblioteca/modelos/remito.py", 100, 119)], "auditoria/diagramas/patron_state.png"),
        ("Strategy: política de multas",
         "La multa por atraso puede calcularse de varias formas (por tipo de material, monto fijo por día, condonación) y la biblioteca la cambia sin cambiar el proceso de devolución.",
         "`EstrategiaMulta` (estrategias_multa.py:21) con `MultaPorDiasDeAtraso` (:34), `MultaPorTipoDeMaterial` (:47) y `SinMulta` (:66). "
         "`CalculadoraDeMultas` (:77) es el contexto con `cambiar_estrategia`. Cliente: `ServicioDevoluciones` (servicio_devoluciones.py:23). "
         "La GUI la cambia en ejecución desde un combo (pestania_devoluciones.py:52), comprobado en CP-25.",
         "Sin el patrón, el servicio tendría un `if politica == ...` y cada política nueva obligaría a modificarlo. A diferencia de State, la estrategia la elige quien configura el contexto, no el objeto.",
         [("biblioteca/patrones/estrategias_multa.py", 21, 32), ("biblioteca/patrones/estrategias_multa.py", 77, 96)],
         "auditoria/diagramas/patron_strategy.png"),
        ("Iterator: recorrido del catálogo de ejemplares",
         "Varios casos de uso recorren ejemplares filtrando por sede, libro, disponibilidad o 'en circulación'.",
         "`CatalogoEjemplares` (iterador_catalogo.py:31) es el agregado y oculta la lista; `IteradorEjemplares` (:11) implementa `__iter__`/`__next__` "
         "y saltea lo que no cumple la condición. Lo usan `ServicioDisponibilidad` (servicio_disponibilidad.py:34) y `ServicioReservas` (servicio_reservas.py:40).",
         "Sin el patrón, cada servicio repetiría los mismos `for` con `if` y dependería de la estructura interna de la colección.",
         [("biblioteca/patrones/iterador_catalogo.py", 11, 52)], "auditoria/diagramas/patron_iterator.png"),
        ("DAO: acceso a datos",
         "Separar las consultas SQLAlchemy de las reglas de negocio y que un caso de uso confirme o deshaga todas sus operaciones juntas.",
         "Un DAO por entidad en `biblioteca/dao/` (por ejemplo `PrestamoDAO`, prestamo_dao.py:6). Reciben la sesión abierta por el servicio con "
         "`obtener_sesion()` (base_de_datos.py:54), que hace commit o rollback.",
         "Sin DAO, los servicios mezclarían consultas con reglas y cada uno manejaría su propia transacción. Es el patrón de la guía 'Objetos de acceso a datos' de la cátedra, sobre ORM en lugar de sqlite3 directo.",
         [("biblioteca/dao/prestamo_dao.py", 6, 20), ("biblioteca/configuracion/base_de_datos.py", 54, 68)],
         "auditoria/diagramas/patron_dao.png"),
    ]
    for titulo, que, donde, porque, fragmentos, diagrama in patrones:
        h(2, titulo)
        tabla(["", ""], [["Qué resuelve", que], ["Dónde", donde], ["Por qué", porque]], "ficha")
        for ruta, desde, hasta in fragmentos:
            codigo(fragmento(ruta, desde, hasta), f"{ruta}:{desde}-{hasta}")
        img(diagrama, f"Diagrama de clases: {titulo.split(':')[0]}", 13)
    h(2, "Usos implícitos, problemas de aplicación y patrones que faltan")
    tabla(["Tipo", "Detalle", "Dónde"], [
        ["Implícito", "Fábrica simple / registro: `crear_estado(nombre)` reconstruye el estado desde el nombre guardado.", "`biblioteca/patrones/estados_remito.py:213`"],
        ["Implícito", "Estrategia con funciones: `EXPORTADORES = {'pdf': ..., 'csv': ...}` elige el exportador por formato.", "`biblioteca/reportes/exportacion.py:7`"],
        ["Implícito", "Unidad de trabajo: el context manager de sesión agrupa las operaciones del caso de uso.", "`biblioteca/configuracion/base_de_datos.py:54`"],
        ["No es Singleton", "El motor es una variable global de módulo (`motor`), no un Singleton de la cátedra. Funciona, pero no conviene presentarlo como patrón.", "`biblioteca/configuracion/base_de_datos.py:12`"],
        ["Mal aplicado (P1)", "Las consecuencias del State se resuelven en el servicio comparando `nombre_estado` con strings.", "`biblioteca/servicios/servicio_remitos.py:43`"],
        ["Resuelto (P2)", "La estrategia vigente se guarda en la tabla `configuracion` y se recupera al iniciar; consola y GUI comparten la política.", "`biblioteca/servicios/servicio_devoluciones.py:18`"],
        ["Falta (P3)", "El ejemplar tiene 7 estados manejados con enum y asignaciones dispersas; es el candidato natural a otro State.", "`biblioteca/modelos/enumeraciones.py:25`"],
        ["Repetición", "Los 8 DAO repiten crear/actualizar/obtener_por_id; un DAO base genérico lo evitaría.", "`biblioteca/dao/*.py`"],
    ])

    # 6 -------------------------------------------------------------------------------------- Lógica
    h(1, "6. Lógica de la biblioteca y de los remitos", "logica")
    h(2, "Entidades")
    tabla(["Entidad", "Identificador de negocio", "Estados", "Regla principal"], [
        ["Sede", "nombre (único, sin distinguir mayúsculas)", "ACTIVA, BAJA", "Baja lógica: conserva historial y no inicia operaciones"],
        ["Socio", "DNI de 7 u 8 dígitos (único)", "HABILITADO, INHABILITADO (con motivo)", "Inhabilitado no recibe préstamos ni reserva"],
        ["Libro", "ISBN de 10 o 13 dígitos (único)", "-", "tipo_material fija la tarifa de multa"],
        ["Ejemplar", "código de inventario (único, mayúsculas)", "DISPONIBLE, RESERVADO, PRESTADO, EN_TRANSITO, EN_REPARACION, EXTRAVIADO, BAJA", "Pertenece siempre a la misma sede aunque viaje"],
        ["Préstamo", "id", "SOLICITADO, EN_ENVIO, LISTO_PARA_RETIRO, ACTIVO, DEVUELTO, CANCELADO", "14 días desde la entrega; 'vencido' es calculado"],
        ["Reserva", "id", "PENDIENTE, DISPONIBLE, CUMPLIDA, CANCELADA, VENCIDA", "Sobre el libro; FIFO; 3 días para retirar"],
        ["Remito", "número R-000001 (único)", "9 estados (State)", "Ida y retorno; cada cambio deja un movimiento"],
        ["Devolución", "uno por préstamo", "multa pagada sí/no", "Multa por atraso (Strategy) + $1000 por daño nuevo"],
    ])
    h(2, "Reglas de negocio verificadas")
    tabla(["Regla", "Código", "Caso E2E"], [
        ["Un ejemplar prestado no se vuelve a prestar", "servicio_prestamos.py:60", "CN-19"],
        ["Un ejemplar en reparación no se presta", "servicio_prestamos.py:60", "CN-44"],
        ["Un ejemplar apartado solo lo retira el titular de la reserva", "servicio_prestamos.py:53", "CP-19, CN-38"],
        ["No se presta a socios inhabilitados, con vencidos o con multas impagas", "servicio_prestamos.py:37-43", "CN-18, CN-22, CN-23"],
        ["Solo se prestan ejemplares de la sede de origen", "servicio_prestamos.py:46", "CN-20 (GUI) y test unitario"],
        ["No se opera con sedes dadas de baja (préstamo, reserva, despacho, alta de ejemplar)", "reglas.py:11, estados_remito.py:83, servicio_catalogo.py:97", "CN-13, CN-21, CN-29, CN-37"],
        ["Solo se reserva si la sede tiene el libro y ninguno disponible", "servicio_reservas.py:40-49", "CN-33, CN-34"],
        ["Se cobra $1000 solo por un daño nuevo; el ejemplar pasa a reparación", "estrategias_multa.py:98, servicio_devoluciones.py:74", "CP-24"],
        ["Una devolución interbibliotecaria genera el remito de retorno", "servicio_devoluciones.py:67", "CP-18"],
    ], "")
    h(2, "Diagramas de estados")
    img("auditoria/diagramas/estados_ejemplar.png", "Ejemplar (estados persistidos como enum; transiciones repartidas en servicios)", 16.5)
    img("auditoria/diagramas/estados_prestamo.png", "Préstamo", 15)
    img("auditoria/diagramas/estados_reserva.png", "Reserva", 12)
    img("auditoria/diagramas/estados_remito.png", "Remito (patrón State)", 16.5)
    h(2, "Remitos")
    tabla(["Tipo (sentido)", "Cuándo se genera", "Origen → destino", "Quién lo genera"], [
        ["IDA", "Al registrar una solicitud de préstamo interbibliotecario", "Sede propietaria → sede de retiro", "`ServicioPrestamos.solicitar_prestamo_interbibliotecario` (servicio_prestamos.py:121)"],
        ["RETORNO por devolución", "Al registrar la devolución de un préstamo interbibliotecario", "Sede de retiro → sede propietaria", "`ServicioDevoluciones.registrar_devolucion` (servicio_devoluciones.py:70)"],
        ["RETORNO por daño", "Cuando un remito de ida llega dañado o se registra daño antes de entregar", "Sede de retiro → sede propietaria", "`_resolver_consecuencias` (servicio_remitos.py:47)"],
    ])
    tabla(["Aspecto", "Detalle"], [
        ["Datos que lleva", "número, préstamo (socio y ejemplar), sede de origen, sede de destino, sentido, estado, fecha de generación y el historial de movimientos (estado anterior, nuevo, fecha y hora, observación: ubicación, motivo o daño)."],
        ["Numeración", "Correlativa con formato R-000001, calculada como max(id)+1 en `RemitoDAO.siguiente_numero` (remito_dao.py:26); columna `numero` UNIQUE."],
        ["Clase que lo genera", "`Remito.generar` (remito.py:55), llamada desde `generar_remito` (servicio_remitos.py:18)."],
        ["Formato de exportación", "**No implementado.** El remito se ve en la pestaña Remitos (con su trayecto) y en el menú de consola 'Ver trazabilidad'. "
                                   "Los únicos archivos exportables son los 5 reportes (PDF/CSV); el reporte 1 lista el material en tránsito y el 5 las incidencias de remitos."],
    ])
    h(2, "Ejemplos reales de cada tipo de remito")
    p("Como no hay exportación del remito, la evidencia es la pantalla real del sistema con los datos de ejemplo (CP-26) y la salida real de consola "
      "(`ejemplos/remitos/trazabilidad_remitos_consola.txt`).")
    galeria([("capturas/CP-26_02_trayecto_R-000001.png", "IDA entregada al socio: R-000001 (muestra el defecto D1: 'Cerrado' con hora anterior a 'Recibido')"),
             ("capturas/CP-26_02_trayecto_R-000002.png", "IDA con daño en la recepción: R-000002"),
             ("capturas/CP-26_02_trayecto_R-000003.png", "RETORNO por daño: R-000003"),
             ("capturas/CP-26_02_trayecto_R-000004.png", "IDA en viaje con punto de control: R-000004")])
    h(2, "Reportes exportados por el sistema")
    nombres = {"activos": "Préstamos activos y material en tránsito", "vencidos": "Préstamos vencidos y multas acumuladas",
               "solicitados": "Libros más solicitados inter-sedes", "movimientos": "Movimiento de ejemplares entre sedes",
               "tasa": "Tasa de envíos exitosos e incidencias"}
    tabla(["Reporte", "PDF", "CSV"], [[n, f"[reporte_{k}.pdf](ejemplos/reportes/reporte_{k}.pdf)", f"[reporte_{k}.csv](ejemplos/reportes/reporte_{k}.csv)"]
                                       for k, n in nombres.items()])
    galeria([(f"ejemplos/reportes/vistas/reporte_{k}.png", f"Primera página del PDF: {n}") for k, n in nombres.items()])

    # 7 -------------------------------------------------------------------------------------- Flujos
    h(1, "7. Flujos positivos y negativos", "flujos")
    p("Cada flujo se ejecutó de verdad sobre la interfaz gráfica (Xvfb, 1320×780) con el driver E2E. Las capturas son de esa ejecución. "
      "Los diálogos no muestran barra de título porque Xvfb corre sin administrador de ventanas; el título ('Listo', 'Operación rechazada', 'Atención', 'Confirmar') lo verifica la prueba.")
    area_actual = None
    for c in CASOS:
        if c.funcionalidad != area_actual:
            area_actual = c.funcionalidad
            h(2, area_actual)
        estado = resultados.get(c.prueba.split("::")[1], ("Sin ejecutar",))[0]
        detalle_inicio(f"{c.id} · {c.descripcion}", estado)
        tabla(["", ""], [
            ["Tipo", "Flujo positivo" if c.tipo == "CP" else "Flujo negativo"],
            ["Precondiciones", "; ".join(c.dado)],
            ["Pasos", "\n".join(f"{i}. {x}" for i, x in enumerate(c.cuando, 1))],
            ["Resultado esperado", "\n".join(c.entonces)],
            ["Resultado de la ejecución", estado],
            ["Requisitos", ", ".join(c.requisitos)],
            ["Prueba", f"`e2e/{c.prueba}`"],
        ], "ficha")
        caps = capturas_de(c.id)
        if caps:
            galeria([(f"capturas/{x.name}", x.stem.split("_", 1)[1].replace("_", " ")) for x in caps])
        elif c.id in ("CP-29", "CN-43"):
            p(f"Salida real de consola: `capturas/{c.id}_salida_consola.txt`.")
        detalle_fin()

    # 8 -------------------------------------------------------------------------------------- Pruebas
    h(1, "8. Plan y resultados de las pruebas E2E", "pruebas")
    h(2, "Enfoque")
    ul([
        "La aplicación es de escritorio (Tkinter), no web, así que Playwright no aplica a la GUI. Se construyó un driver (`e2e/driver.py`) que maneja la aplicación real: "
        "hace clic en los botones, completa los formularios modales, lee el texto de los cuadros de mensaje de Tk y los responde. No reemplaza ninguna función de la app "
        "(los tests de `tests/test_interfaz_grafica.py` sí simulan los diálogos).",
        "Cada prueba usa una base SQLite nueva en un archivo temporal; las precondiciones se cargan con los servicios del sistema y los datos son fijos (fechas relativas a hoy), por lo que es reproducible.",
        "La consola se prueba ejecutando `python main.py` como proceso con entradas por stdin.",
        "Casos en Gherkin con ID: `e2e/features/*.feature` (en español) y la planilla `docs/pruebas/casos_xray.csv` (Summary | Description | Repositorio | Tipo de test | Gherkin) para Jira/Xray.",
    ])
    h(2, "Resultados")
    tabla(["Suite", "Resultado", "Detalle"], [
        ["E2E (`e2e/`)", f"{estados_e2e['Pasa']} pasan · {estados_e2e['Falla esperada (defecto)']} falla esperada · {estados_e2e['Falla']} fallan",
         "CN-45 documentaba el defecto D2 como falla esperada; desde su corrección pasa. Log: `docs/pruebas/resultados_e2e.txt`, JUnit: `resultados_e2e.xml`."],
        ["Unitarios (`tests/`)", unit, "El test del defecto D1 solo falla si se ejecuta antes de las 10:00 (depende de la hora del reloj). Log: `docs/pruebas/resultados_unitarios.txt`."],
        ["Cobertura unitaria", "97 % de `biblioteca` sin las interfaces", "Medido con coverage 7 (el README declaraba 96 %; se actualizó)."],
    ])
    por_area = OrderedDict()
    for c in CASOS:
        por_area.setdefault(c.funcionalidad, Counter())[resultados.get(c.prueba.split("::")[1], ("Sin ejecutar",))[0]] += 1
        por_area[c.funcionalidad][c.tipo] += 1
    tabla(["Funcionalidad", "CP", "CN", "Pasan", "Falla esperada", "Fallan", "Feature"],
          [[a, n["CP"], n["CN"], n["Pasa"], n["Falla esperada (defecto)"], n["Falla"], f"`{features[a]}`"] for a, n in por_area.items()])
    h(2, "Casos y resultado")
    tabla(["ID", "Summary (Xray)", "Resultado", "Tiempo"],
          [[c.id, c.summary, resultados.get(c.prueba.split("::")[1], ("Sin ejecutar", 0))[0],
            f"{resultados.get(c.prueba.split('::')[1], ('', 0))[1]:.2f} s"] for c in CASOS], "casos")
    h(2, "Ejemplo de caso en Gherkin")
    ejemplo = next(c for c in CASOS if c.id == "CN-27")
    codigo("\n".join(["# language: es", f"Escenario: {ejemplo.id} {ejemplo.descripcion}"] +
                     [f"  {pal if i == 0 else 'Y'} {x}" for pal, xs in (("Dado", ejemplo.dado), ("Cuando", ejemplo.cuando), ("Entonces", ejemplo.entonces))
                      for i, x in enumerate(xs)]), features[ejemplo.funcionalidad])
    h(2, "Trazabilidad requisito → casos de prueba")
    tabla(["Req.", "Descripción", "Estado", "Casos E2E", "Tests unitarios adicionales"],
          [[r.id, r.descripcion, r.estado, ", ".join(casos_por_req[r.id]) or "-", "; ".join(r.unitarios) or "-"] for r in REQUISITOS])

    # 9 -------------------------------------------------------------------------------------- Faltantes
    h(1, "9. Faltantes y mejoras sugeridas", "faltantes")
    for x in pendientes:
        h(3, f"{x.id} · {x.titulo} (prioridad {x.prioridad.lower()})")
        p(x.detalle)
        p("**Evidencia:** " + ", ".join(f"`{e}`" for e in x.evidencia))
        p("**Arreglo sugerido:** " + x.arreglo)
    if resueltos:
        h(3, "Hallazgos corregidos después de la auditoría")
        tabla(["#", "Hallazgo original", "Qué se cambió"],
              [[x.id, f"**{x.titulo}.** {x.detalle}", x.resuelto] for x in resueltos])
    h(3, "Opcionales de la consigna")
    tabla(["Req.", "Proceso", "Estado", "Observaciones"],
          [[r.id, r.descripcion, r.estado, r.observaciones or "-"] for r in REQUISITOS if r.grupo == "Proceso opcional"])

    # 10 ------------------------------------------------------------------------------------- Anexo
    h(1, "10. Anexo: cómo ejecutar el sistema y las pruebas", "anexo")
    codigo("""git clone https://github.com/LucasChas/TP-DAO-G054.git && cd TP-DAO-G054
git checkout desarrollo            # o la rama de esta auditoría: claude/eager-wright-hnv32r
python3.12 -m venv .venv && source .venv/bin/activate      # Windows: .venv\\Scripts\\activate
pip install -r requirements.txt
# Linux: Tkinter va aparte -> sudo apt install python3-tk (debe ser del mismo Python del venv)

python main.py              # consola  (opción 10: cargar datos de ejemplo)
python main.py --grafica    # interfaz gráfica (Archivo > Cargar datos de ejemplo)""", "Sistema")
    codigo("""python -m pytest tests -q                                   # 247 tests unitarios
xvfb-run -a -s "-screen 0 1320x780x24" python -m pytest e2e -v   # E2E (en Windows/macOS sin xvfb-run)

# Regenerar capturas, reportes de ejemplo y documentación:
E2E_CAPTURAS=$PWD/docs/capturas E2E_EJEMPLOS=$PWD/docs/ejemplos/reportes \\
  xvfb-run -a -s "-screen 0 1320x780x24" python -m pytest e2e -v --junitxml=docs/pruebas/resultados_e2e.xml
PYTHONPATH=. python docs/auditoria/construir_documentacion.py""", "Pruebas y documentación")
    nota("Entorno de la auditoría: Linux, Python 3.12.3 (venv), Tk 8.6, Xvfb. Python 3.13 del sistema no trae Tkinter; se usó 3.12, que cumple el requisito '3.12 o superior' del README.")
    return resumen


# --------------------------------------------------------------------------------- HTML


def inline_html(texto):
    t = html.escape(str(texto))
    t = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r'<a href="\2">\1</a>', t)
    t = re.sub(r"`([^`]+)`", r"<code>\1</code>", t)
    t = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", t)
    return t.replace("\n", "<br>")


def chip(estado):
    clase = {"Cumple": "ok", "Pasa": "ok", "Parcial": "warn", "Falla esperada (defecto)": "warn",
             "No cumple": "bad", "Falla": "bad", "No implementado": "off", "Alta": "bad", "Media": "warn", "Baja": "off"}.get(estado)
    return f'<span class="sello {clase}">{html.escape(estado)}</span>' if clase else None


def render_html(resumen):
    partes, toc, abierta = [], [], False
    for b in B:
        tipo = b[0]
        if tipo == "h":
            _, nivel, texto, ancla = b
            if nivel == 1:
                if abierta:
                    partes.append("</div></details></section>")
                abierta = True
                toc.append((ancla, texto))
                partes.append(f'<section id="{ancla}"><details open><summary><h2>{html.escape(texto)}</h2></summary><div class="cuerpo">')
            else:
                partes.append(f"<h{nivel + 1}>{inline_html(texto)}</h{nivel + 1}>")
        elif tipo == "p":
            partes.append(f"<p>{inline_html(b[1])}</p>")
        elif tipo == "ul":
            partes.append("<ul>" + "".join(f"<li>{inline_html(x)}</li>" for x in b[1]) + "</ul>")
        elif tipo == "table":
            _, enc, filas, clase = b
            cab = "" if not any(enc) else "<thead><tr>" + "".join(f"<th>{html.escape(x)}</th>" for x in enc) + "</tr></thead>"
            cuerpo = "".join("<tr>" + "".join(f"<td>{chip(str(v)) or inline_html(v)}</td>" for v in fila) + "</tr>" for fila in filas)
            partes.append(f'<div class="tabla {clase}"><table>{cab}<tbody>{cuerpo}</tbody></table></div>')
        elif tipo == "img":
            _, ruta, pie, _ = b
            partes.append(f'<figure class="diagrama"><a href="{ruta}" class="ampliar"><img src="{ruta}" alt="{html.escape(pie)}" loading="lazy"></a>'
                          f"<figcaption>{html.escape(pie)}</figcaption></figure>")
        elif tipo == "code":
            partes.append(f'<figure class="codigo"><figcaption>{html.escape(b[2])}</figcaption><pre><code>{html.escape(b[1])}</code></pre></figure>')
        elif tipo == "note":
            partes.append(f'<div class="nota {b[2]}">{inline_html(b[1])}</div>')
        elif tipo == "gallery":
            partes.append('<div class="galeria">' + "".join(
                f'<figure><a href="{r}" class="ampliar"><img src="{r}" alt="{html.escape(pie)}" loading="lazy"></a><figcaption>{html.escape(pie)}</figcaption></figure>'
                for r, pie in b[1]) + "</div>")
        elif tipo == "details_start":
            partes.append(f'<details class="caso"><summary><span>{html.escape(b[1])}</span>{chip(b[2]) or ""}</summary><div>')
        elif tipo == "details_end":
            partes.append("</div></details>")
    partes.append("</div></details></section>")
    e2e = resumen["e2e"]
    tarjetas = [
        (f"{resumen['porcentaje']:.1f} %", "cumplimiento de lo obligatorio", f"{resumen['cuenta']['Cumple']} de {resumen['total']} cumplen · {resumen['cuenta']['Parcial']} parciales"),
        (f"{e2e['Pasa']}/{sum(e2e.values())}", "casos E2E que pasan",
         f"{e2e['Falla esperada (defecto)']} defecto documentado (xfail)" if e2e['Falla esperada (defecto)'] else "D2, P2 y C1 corregidos"),
        (resumen["unit_corto"], "tests unitarios que pasan", "D1 falla antes de las 10:00 (hora del reloj)"),
        ("4", "patrones aplicados", "State · Strategy · Iterator · DAO"),
    ]
    plantilla = (Path(__file__).parent / "plantilla_informe.html").read_text(encoding="utf-8")
    return (plantilla.replace("{{FECHA}}", HOY)
            .replace("{{TOC}}", "".join(f'<li><a href="#{a}">{html.escape(t)}</a></li>' for a, t in toc))
            .replace("{{TARJETAS}}", "".join(f'<div class="dato"><strong>{a}</strong><span>{b}</span><small>{c}</small></div>' for a, b, c in tarjetas))
            .replace("{{CONTENIDO}}", "\n".join(partes)))


# --------------------------------------------------------------------------------- DOCX


def render_docx(resumen, destino):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document(DOCS / "plantilla" / "Presentacion-grupo.docx")
    from PIL import Image

    def ancho_imagen(ruta, ancho_cm, alto_max_cm=21.0):
        """Ancho en cm que respeta el ancho pedido sin pasar el alto de una página."""
        with Image.open(DOCS / ruta) as im:
            w, h_px = im.size
        return Cm(min(ancho_cm, alto_max_cm * w / h_px))
    cuerpo = doc.element.body

    def runs(parrafo, texto, tam=None):
        for trozo in re.split(r"(`[^`]+`|\*\*[^*]+\*\*|\[[^\]]+\]\([^)]+\))", str(texto)):
            if not trozo:
                continue
            if trozo.startswith("`"):
                r = parrafo.add_run(trozo[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(8.5)
            elif trozo.startswith("**"):
                r = parrafo.add_run(trozo[2:-2]); r.bold = True
            elif trozo.startswith("["):
                m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", trozo); r = parrafo.add_run(f"{m.group(1)} (docs/{m.group(2)})")
            else:
                r = parrafo.add_run(trozo)
            if tam and not trozo.startswith("`"):
                r.font.size = Pt(tam)

    def sombrear(celda, color):
        tc = celda._tc.get_or_add_tcPr(); s = OxmlElement("w:shd")
        s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), color)
        tc.insert_element_before(s, "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark")

    def agregar_tabla(contenedor, enc, filas, anchos=None, tam=8):
        t = contenedor.add_table(rows=0, cols=len(enc))
        t.style = doc.styles["Table Grid"] if "Table Grid" in [s.name for s in doc.styles] else None
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        if any(enc):
            fila = t.add_row().cells
            for i, x in enumerate(enc):
                fila[i].text = ""; runs(fila[i].paragraphs[0], x, tam); fila[i].paragraphs[0].runs[0].bold = True
                sombrear(fila[i], "DCE6F0")
        for f in filas:
            celdas = t.add_row().cells
            for i, v in enumerate(f):
                lineas = str(v).split("\n")
                celdas[i].text = ""
                runs(celdas[i].paragraphs[0], lineas[0], tam)
                for extra in lineas[1:]:
                    runs(celdas[i].add_paragraph(), extra, tam)
                color = {"Cumple": "D8EFD9", "Pasa": "D8EFD9", "Parcial": "FCEFC7", "Falla esperada (defecto)": "FCEFC7",
                         "Falla": "F6D5D3", "No cumple": "F6D5D3", "Alta": "F6D5D3", "Media": "FCEFC7"}.get(str(v))
                if color:
                    sombrear(celdas[i], color)
        bordes(t)
        total = sum(anchos) if anchos else None
        if anchos:
            t.autofit = False
            for i, columna in enumerate(t.columns):
                columna.width = Cm(anchos[i] * 17 / total)
            for fila in t.rows:
                for i, c in enumerate(fila.cells):
                    c.width = Cm(anchos[i] * 17 / total)
        return t

    def bordes(t):
        tbl_pr = t._tbl.tblPr
        b = OxmlElement("w:tblBorders")
        for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{lado}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "8A9BB0"); b.append(e)
        # El esquema fija el orden de los hijos de tblPr: tblBorders va antes de shd, tblLayout, tblCellMar y tblLook.
        tbl_pr.insert_element_before(b, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription")

    # Sección 1: tabla del grupo de la plantilla.
    t_grupo = doc.tables[0]
    p0 = t_grupo.rows[0].cells[0].paragraphs[0]
    for r in p0.runs[1:]:
        r.text = ""
    p0.runs[0].text = "Número de grupo: G054        Tema: 1 – Biblioteca: Sistema de Préstamos Interbibliotecarios"
    anidada = t_grupo.rows[2].cells[0].tables[0]
    for fila, (legajo, nombre) in zip(anidada.rows[1:], INTEGRANTES):
        fila.cells[0].text = legajo
        fila.cells[1].text = nombre

    # Secciones 2: objetivos, alcances y diagramas dentro de los recuadros de la plantilla.
    def bloques_entre(inicio, fin):
        dentro, salida = False, []
        for b in B:
            if b[0] == "h" and b[2] == inicio:
                dentro = True; continue
            if b[0] == "h" and (b[2] == fin or (b[1] == 1 and dentro)):
                if dentro:
                    break
            if dentro:
                salida.append(b)
        return salida

    def volcar(celda, bloques):
        celda.paragraphs[0].text = ""
        primero = True
        for b in bloques:
            if b[0] == "ul":
                for x in b[1]:
                    par = celda.paragraphs[0] if primero else celda.add_paragraph()
                    primero = False
                    runs(par, "• " + x, 9.5)
            elif b[0] == "p":
                par = celda.paragraphs[0] if primero else celda.add_paragraph(); primero = False
                runs(par, b[1], 9)
            elif b[0] == "table":
                agregar_tabla(celda, b[1], b[2], tam=8.5)
            elif b[0] == "img":
                par = celda.paragraphs[0] if primero else celda.add_paragraph(); primero = False
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                par.add_run().add_picture(str(DOCS / b[1]), width=ancho_imagen(b[1], min(b[3], 16.5), 19))
                pie = celda.add_paragraph(); runs(pie, b[2], 8); pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif b[0] == "h":
                par = celda.add_paragraph(); runs(par, f"**{b[2]}**", 10)

    volcar(doc.tables[1].rows[0].cells[0], bloques_entre("Objetivos del sistema", "Alcances"))
    volcar(doc.tables[2].rows[0].cells[0], bloques_entre("Alcances", "DER (esquema real de la base)"))
    der = [b for b in B if b[0] in ("img", "p", "h") and B.index(b) > B.index(next(x for x in B if x[0] == "h" and x[2] == "DER (esquema real de la base)"))
           and B.index(b) < B.index(next(x for x in B if x[0] == "h" and x[2] == "3. Arquitectura"))]
    volcar(doc.tables[3].rows[0].cells[0], [b for b in der if b[0] != "h" or b[2] != "DER (esquema real de la base)"])

    # El último Heading 2 vacío de la plantilla se reutiliza como punto de inserción.
    ancla = doc.paragraphs[-1]
    ancla.text = ""

    def nuevo(estilo=None):
        par = doc.add_paragraph(style=estilo)
        return par

    en_secciones_propias = False
    for b in B:
        if b[0] == "h" and b[1] == 1 and b[2].startswith("3."):
            en_secciones_propias = True
        if not en_secciones_propias:
            continue
        tipo = b[0]
        if tipo == "h":
            estilo = {1: "Heading 2", 2: "Heading 3", 3: "Heading 3"}[b[1]]
            if b[1] == 1:
                doc.add_page_break()
            runs(nuevo(estilo), b[2])
        elif tipo == "p":
            runs(nuevo(), b[1], 10)
        elif tipo == "ul":
            for x in b[1]:
                par = nuevo("List Bullet") if "List Bullet" in [s.name for s in doc.styles] else nuevo()
                runs(par, ("" if par.style.name == "List Bullet" else "• ") + x, 10)
        elif tipo == "table":
            _, enc, filas, clase = b
            anchos = {"matriz": [1.0, 3.0, 1.4, 4.6, 2.6, 4.4], "ficha": [3, 14], "casos": [1.5, 10.5, 3.2, 1.8]}.get(clase)
            if clase == "matriz":
                filas = [[str(v).replace("`", "") for v in f] for f in filas]
            agregar_tabla(doc, enc, filas, anchos, 7.5 if clase == "matriz" else 8.5)
            nuevo()
        elif tipo == "img":
            par = nuevo(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.add_run().add_picture(str(DOCS / b[1]), width=ancho_imagen(b[1], b[3]))
            pie = nuevo(); runs(pie, b[2], 8.5); pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif tipo == "code":
            pie = nuevo(); runs(pie, f"**{b[2]}**", 8.5)
            for linea in b[1].split("\n"):
                par = nuevo(); r = par.add_run(linea or " "); r.font.name = "Consolas"; r.font.size = Pt(7.5)
                par.paragraph_format.space_after = Pt(0); par.paragraph_format.space_before = Pt(0)
                sombra = OxmlElement("w:shd"); sombra.set(qn("w:val"), "clear"); sombra.set(qn("w:fill"), "F1F4F8")
                par._p.get_or_add_pPr().insert_element_before(
                    sombra, "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct", "w:topLinePunct",
                    "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
                    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
                    "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")
            nuevo()
        elif tipo == "note":
            par = nuevo(); r = par.add_run("Nota: "); r.bold = True
            r.font.color.rgb = RGBColor(0x9A, 0x5B, 0x00) if b[2] == "alerta" else RGBColor(0x24, 0x52, 0x7A)
            runs(par, b[1], 9.5)
        elif tipo == "gallery":
            for ruta, pie in b[1]:
                par = nuevo(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                par.add_run().add_picture(str(DOCS / ruta), width=ancho_imagen(ruta, 13))
                q = nuevo(); runs(q, pie, 8); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif tipo == "details_start":
            runs(nuevo("Heading 3"), f"{b[1]} — {b[2]}")
    # Mover el ancla vacía al final no hace falta: queda como separador antes de la sección 3.
    doc.core_properties.title = "TP Integrador DAO – G054 – Presentación y auditoría"
    doc.core_properties.author = "Grupo G054"
    doc.save(destino)
    return destino


# --------------------------------------------------------------------------------- principal


def main():
    rotas = verificar_referencias()
    if rotas:
        print("Referencias inexistentes:", rotas)
        sys.exit(1)
    resultados = resultados_junit()
    sin_resultado = [c.id for c in CASOS if c.prueba.split("::")[1] not in resultados]
    if sin_resultado:
        print("Casos sin resultado en el JUnit:", sin_resultado)
        sys.exit(1)
    features = escribir_features()
    xray = escribir_xray()
    resumen = construir_contenido(resultados, features)
    (DOCS / "informe_auditoria.html").write_text(render_html(resumen), encoding="utf-8")
    docx = render_docx(resumen, DOCS / "Presentacion-grupo-G054.docx")
    print(f"Cumplimiento: {resumen['porcentaje']:.1f} % · E2E: {dict(resumen['e2e'])} · Unitarios: {resumen['unit']}")
    print("Generado:", ", ".join(str(x.relative_to(RAIZ)) for x in [xray, DOCS / "informe_auditoria.html", docx]),
          "y", len(features), "features")


if __name__ == "__main__":
    main()
